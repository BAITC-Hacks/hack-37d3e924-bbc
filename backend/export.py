from io import BytesIO
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn


def export_docx(meeting):
    result = meeting['result']
    document = Document()
    for style in document.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        document.styles[name].font.name = 'Arial'
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    document.styles['Normal'].font.size = Pt(10)
    document.styles['Title'].font.size = Pt(20)
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    document.core_properties.author = ''
    document.core_properties.last_modified_by = '' 
    if meeting['mode'] == 'fixture':
        document.add_paragraph('ТЕСТОВЫЙ РЕЗУЛЬТАТ — синтетический пример; модели не запускались.')
    document.add_heading(meeting['title'], 0)
    document.add_paragraph(f"{meeting['meeting_datetime']} • {meeting['timezone']} • Версия {meeting['revision']}")
    document.add_heading('Участники', 1)
    names = {p['id']: p['name'] or 'требует уточнения' for p in result['participants']}
    for participant in result['participants']:
        document.add_paragraph(f"{names[participant['id']]} ({participant['id']}); голоса: " +
                               (', '.join(participant['speaker_ids']) or 'требует уточнения'))
    document.add_heading('Краткое содержание', 1)
    document.add_paragraph(result['summary'])
    document.add_heading('Поручения', 1)
    for task in result['tasks']:
        document.add_paragraph(task['text'], style='List Number')
        document.add_paragraph(f"Исполнитель: {names.get(task['assignee_id'], 'требует уточнения')}. "
                               f"Срок: {task['due_date'] or 'требует уточнения'}. "
                               f"Источники: {', '.join(task['source_segment_ids'])}. "
                               + ('Требует проверки.' if task['needs_review'] else ''))
    document.add_heading('Транскрипт', 1)
    for segment in result['segments']:
        document.add_paragraph(f"[{segment['start']:.1f}–{segment['end']:.1f}] "
                               f"{segment['speaker_id']} ({segment['id']}): {segment['text']}")
    if result['warnings']:
        document.add_heading('Предупреждения', 1)
        for warning in result['warnings']:
            document.add_paragraph(warning)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
