from io import BytesIO
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from core import stamp

def build_docx(title, meeting_date, analysis, segments, names, include_transcript=True):
    doc = Document()
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for style in ['Normal','Title','Heading 1','Heading 2']:
        doc.styles[style].font.name = 'Arial'
        doc.styles[style].font.color.rgb = RGBColor(0,0,0)
    doc.styles['Normal'].font.size = Pt(10)
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Heading 1'].font.size = Pt(14)
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.add_paragraph('Протокол совещания', 'Title')
    doc.add_paragraph(title or 'Совещание')
    doc.add_paragraph('Дата: ' + (meeting_date or 'не указана'))
    doc.add_paragraph('Черновик для проверки секретарём. Имена, сроки и факты требуют подтверждения.')
    doc.add_heading('Краткое содержание', 1)
    doc.add_paragraph(analysis.get('summary') or 'Не сформировано')
    doc.add_heading('Поручения', 1)
    tasks = analysis.get('tasks', [])
    if tasks:
        table = doc.add_table(rows=1, cols=4)
        table.autofit = False
        widths = [1.0, 7.6, 4.2, 4.2]
        for col, width in zip(table.columns, widths):
            col.width = Cm(width)
        for cell, text in zip(table.rows[0].cells, ['№','Поручение','Ответственный','Срок']):
            cell.text = text
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7EDF1'); cell._tc.get_or_add_tcPr().append(shade)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for index, task in enumerate(tasks, 1):
            cells = table.add_row().cells
            values = [str(index), task.get('task',''), task.get('owner') or 'Не указан',
                      task.get('due_date') or task.get('due_text') or 'Не указан']
            for cell, value in zip(cells,values):
                cell.text = str(value)
        for row in table.rows:
            for cell, width in zip(row.cells,widths):
                cell.width = Cm(width)
                pr = cell._tc.get_or_add_tcPr()
                borders = OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    el = OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
                pr.append(borders)
                margins = OxmlElement('w:tcMar')
                for edge in ['top','left','bottom','right']:
                    el = OxmlElement('w:'+edge); el.set(qn('w:w'),'90'); el.set(qn('w:type'),'dxa'); margins.append(el)
                pr.append(margins)
    else:
        doc.add_paragraph('Подтверждённые поручения не найдены. Проверьте транскрипт.')
    if tasks:
        doc.add_heading('Основания и вопросы для проверки',1)
        for i, task in enumerate(tasks,1):
            when = f"{stamp(task['start'])} — " if task.get('start') is not None else ''
            doc.add_paragraph(f"{i}. {when}{task.get('quote') or 'Добавлено вручную'}")
            if task.get('review'):
                doc.add_paragraph('Проверить: '+task['review'])
    if include_transcript:
        doc.add_page_break()
        doc.add_heading('Транскрипт',1)
        for s in segments:
            p = doc.add_paragraph()
            if s['start'] is not None or s['speaker'] != 'UNKNOWN':
                speaker = names.get(s['speaker'],s['speaker'])
                when = f"{stamp(s['start'])} · " if s['start'] is not None else ''
                p.add_run(f"{when}{speaker}\n").bold = True
            p.add_run(s['text'])
    out = BytesIO(); doc.save(out)
    return out.getvalue()
