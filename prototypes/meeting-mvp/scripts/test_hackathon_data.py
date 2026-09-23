"""Local smoke test for bundled hackathon data.

The report intentionally contains only metadata, counts, statuses and timings.
It must not print or persist extracted meeting text outside .local/data-smoke.
"""
from __future__ import annotations

import json
import os
import hashlib
import sys
import tempfile
import time
from datetime import datetime, timezone
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = ROOT.parents[1]
DATA_ROOT = REPO_ROOT / 'data'
SMOKE_ROOT = ROOT / '.local' / 'data-smoke'

sys.path.insert(0, str(ROOT))

os.environ.setdefault('HF_HUB_OFFLINE', '1')
os.environ.setdefault('TRANSFORMERS_OFFLINE', '1')
os.environ.setdefault('HF_HUB_DISABLE_TELEMETRY', '1')
os.environ.setdefault('DO_NOT_TRACK', '1')

from config import MODEL_DIR  # noqa: E402
from core import read_json, write_json  # noqa: E402
from engine import decode_audio, diarize  # noqa: E402


def _deny_network(event, args):
    if event == 'socket.getaddrinfo':
        raise PermissionError('Network access is disabled for the data smoke test.')
    if event in ('socket.connect', 'socket.sendto') and args and getattr(args[0], 'family', None) in (2, 10):
        raise PermissionError('Network access is disabled for the data smoke test.')


sys.addaudithook(_deny_network)


def _audio_status():
    asr_ready = all((MODEL_DIR / p).is_file() for p in ['asr/model.pt', 'asr/tokens.lst'])
    llm_ready = _mlx_ready()
    return {
        'asr': {
            'status': 'READY_NOT_RUN' if asr_ready else 'BLOCKED',
            'reason': 'ASR is intentionally not run by this smoke test.' if asr_ready else 'Missing source ASR weights.',
        },
        'llm': {
            'status': 'READY_NOT_RUN' if llm_ready else 'BLOCKED',
            'reason': 'LLM is intentionally not run by this smoke test.' if llm_ready else 'Missing MLX runtime or source LLM weights.',
        },
    }


def _mlx_ready():
    from importlib.util import find_spec

    return find_spec('mlx_lm') is not None and all(
        (MODEL_DIR / p).is_file() for p in ['llm/model.safetensors', 'llm/config.json', 'llm/tokenizer.json']
    )


def _diarization_ready():
    required = ['diarization/segmentation.onnx', 'diarization/embedding.onnx']
    missing = [p for p in required if not (MODEL_DIR / p).is_file()]
    return missing == [], missing


def _decode_smoke(mp3_path, run_dir):
    run_dir.mkdir(parents=True, exist_ok=True)
    audio_wav = run_dir / 'audio.wav'
    started = time.perf_counter()
    duration = decode_audio(mp3_path, audio_wav)
    elapsed = time.perf_counter() - started

    import soundfile as sf

    info = sf.info(str(audio_wav))
    checks = {
        'is_16khz': int(info.samplerate) == 16000,
        'is_mono': int(info.channels) == 1,
        'duration_positive': float(duration) > 0,
    }
    return {
        'status': 'PASS' if all(checks.values()) else 'FAIL',
        'elapsed_seconds': round(elapsed, 3),
        'duration_seconds': round(float(duration), 3),
        'wav_sample_rate': int(info.samplerate),
        'wav_channels': int(info.channels),
        'wav_frames': int(info.frames),
        'checks': checks,
    }


def _diarization_smoke(run_dir, duration):
    ready, missing = _diarization_ready()
    if not ready:
        return {'status': 'BLOCKED', 'reason': 'Missing source diarization weights.', 'missing': missing}

    write_json(run_dir / 'request.json', {'num_speakers': -1})
    write_json(run_dir / 'status.json', {'state': 'running', 'label': 'smoke', 'progress': 0})
    started = time.perf_counter()
    try:
        diarize(run_dir)
    except Exception as error:  # Keep sanitized technical metadata, not meeting content.
        return {
            'status': 'FAIL',
            'elapsed_seconds': round(time.perf_counter() - started, 3),
            'error_type': type(error).__name__,
            'reason': 'Diarization failed during local smoke test.',
        }

    turns = read_json(run_dir / 'turns.json')
    speakers = sorted({turn.get('speaker') for turn in turns})
    invalid_turns = [
        index for index, turn in enumerate(turns)
        if not _valid_turn(turn, duration)
    ]
    starts_sorted = all(turns[i]['start'] <= turns[i + 1]['start'] for i in range(max(0, len(turns) - 1)))
    return {
        'status': 'PASS' if turns and not invalid_turns and starts_sorted else 'FAIL',
        'elapsed_seconds': round(time.perf_counter() - started, 3),
        'turn_count': len(turns),
        'speaker_count': len(speakers),
        'checks': {
            'has_turns': bool(turns),
            'starts_sorted': starts_sorted,
            'all_turns_inside_audio': not invalid_turns,
        },
        'invalid_turn_indexes': invalid_turns[:20],
    }


def _valid_turn(turn, duration):
    start = turn.get('start')
    end = turn.get('end')
    return (
        isinstance(start, (int, float))
        and isinstance(end, (int, float))
        and 0 <= start < end
        and end <= duration + 0.25
    )


def _docx_smoke(docx_path, output_dir):
    from docx import Document

    started = time.perf_counter()
    doc = Document(str(docx_path))
    paragraph_texts = [paragraph.text for paragraph in doc.paragraphs]
    cell_texts = [
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    ]
    paragraph_count = len(doc.paragraphs)
    table_count = len(doc.tables)
    nonempty_paragraph_count = sum(1 for paragraph in doc.paragraphs if paragraph.text.strip())
    nonempty_cell_count = sum(
        1
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
        if cell.text.strip()
    )
    roundtrip_path = output_dir / (docx_path.stem + '.roundtrip.docx')
    doc.save(str(roundtrip_path))
    roundtrip = Document(str(roundtrip_path))
    roundtrip_paragraph_texts = [paragraph.text for paragraph in roundtrip.paragraphs]
    roundtrip_cell_texts = [
        cell.text
        for table in roundtrip.tables
        for row in table.rows
        for cell in row.cells
    ]
    checks = {
        'paragraph_texts_equal': paragraph_texts == roundtrip_paragraph_texts,
        'cell_texts_equal': cell_texts == roundtrip_cell_texts,
        'has_readable_content': bool(nonempty_paragraph_count or nonempty_cell_count),
    }
    return {
        'status': 'PASS' if all(checks.values()) else 'FAIL',
        'elapsed_seconds': round(time.perf_counter() - started, 3),
        'paragraph_count': paragraph_count,
        'nonempty_paragraph_count': nonempty_paragraph_count,
        'table_count': table_count,
        'nonempty_cell_count': nonempty_cell_count,
        'roundtrip_file': str(roundtrip_path.relative_to(ROOT)),
        'checks': checks,
    }


def _file_metadata(path):
    stat = path.stat()
    return {
        'file': path.name,
        'bytes': stat.st_size,
        'mtime_utc': datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(timespec='seconds'),
    }


def _sha256(path):
    digest = hashlib.sha256()
    with path.open('rb') as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def _source_snapshot(path):
    return {**_file_metadata(path), 'sha256': _sha256(path)}


def _source_unchanged(before, path):
    after = _source_snapshot(path)
    return after, before['sha256'] == after['sha256'] and before['bytes'] == after['bytes']


def _overall_status(report):
    if not report['audio_files'] or not report['docx_files']:
        return 'FAIL'
    statuses = []
    for item in report['audio_files']:
        statuses.append(item.get('decode', {}).get('status'))
        statuses.append(item.get('diarization', {}).get('status'))
        if not item.get('source_unchanged'):
            statuses.append('FAIL')
    for item in report['docx_files']:
        statuses.append(item.get('docx', {}).get('status'))
        if not item.get('source_unchanged'):
            statuses.append('FAIL')
    if any(status == 'FAIL' for status in statuses):
        return 'FAIL'
    if report['asr']['status'] != 'READY_NOT_RUN' or report['llm']['status'] != 'READY_NOT_RUN':
        return 'PARTIAL'
    return 'PARTIAL'


def main():
    SMOKE_ROOT.mkdir(parents=True, exist_ok=True)
    run_root = Path(tempfile.mkdtemp(prefix='run-', dir=SMOKE_ROOT))
    local_status = _audio_status()
    report = {
        'schema_version': 1,
        'generated_at_utc': datetime.now(timezone.utc).isoformat(timespec='seconds'),
        'data_dir': str(DATA_ROOT.relative_to(REPO_ROOT)),
        'output_dir': str(run_root.relative_to(ROOT)),
        'network': 'disabled_by_audit_hook',
        'audio_files': [],
        'docx_files': [],
        'asr': local_status['asr'],
        'llm': local_status['llm'],
    }

    for index, mp3_path in enumerate(sorted(DATA_ROOT.glob('*.mp3')), 1):
        run_dir = run_root / f'audio-{index:02d}'
        before = _source_snapshot(mp3_path)
        item = {'source_before': before}
        try:
            item['decode'] = _decode_smoke(mp3_path, run_dir)
            item['diarization'] = (
                _diarization_smoke(run_dir, item['decode']['duration_seconds'])
                if item['decode']['status'] == 'PASS'
                else {'status': 'BLOCKED', 'reason': 'Audio decode did not pass.'}
            )
        except Exception as error:
            item['decode'] = {
                'status': 'FAIL',
                'error_type': type(error).__name__,
                'reason': 'Audio decode failed during local smoke test.',
            }
            item['diarization'] = {'status': 'BLOCKED', 'reason': 'Audio decode did not pass.'}
        item['source_after'], item['source_unchanged'] = _source_unchanged(before, mp3_path)
        report['audio_files'].append(item)

    docx_output = run_root / 'docx-roundtrip'
    docx_output.mkdir(parents=True, exist_ok=True)
    for docx_path in sorted(DATA_ROOT.glob('*.docx')):
        before = _source_snapshot(docx_path)
        item = {'source_before': before}
        try:
            item['docx'] = _docx_smoke(docx_path, docx_output)
        except Exception as error:
            item['docx'] = {
                'status': 'FAIL',
                'error_type': type(error).__name__,
                'reason': 'DOCX roundtrip failed during local smoke test.',
            }
        item['source_after'], item['source_unchanged'] = _source_unchanged(before, docx_path)
        report['docx_files'].append(item)

    report['status'] = _overall_status(report)
    report_path = run_root / 'report.json'
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({
        'status': report['status'],
        'report': str(report_path),
        'audio_count': len(report['audio_files']),
        'docx_count': len(report['docx_files']),
    }, ensure_ascii=False))
    raise SystemExit(1 if report['status'] == 'FAIL' else 0)


if __name__ == '__main__':
    main()
