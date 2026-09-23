"""Synthetic DOCX export coverage only."""
from io import BytesIO
import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from export_docx import build_docx


SEGMENTS = [{'id':'S1','speaker':'A','start':None,'end':None,'text':'Синтетикалық тапсырма мәтіні.'}]
NAMES = {'A':'Айжан Қали'}
ANALYSIS = {'summary':'Синтетическое саммари','tasks':[{'task':'Проверить план','owner':'Айжан Қали',
    'due_text':'','due_date':None,'review':'','status':None,'quote':'Поддельная цитата','source_ids':['S1']}],
    'warnings':['Синтетическое предупреждение']}


def docx_text(blob):
    document = Document(BytesIO(blob))
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    return '\n'.join(texts)


def test_docx_includes_participants_without_transcript_and_marks_synthetic_source():
    text = docx_text(build_docx('Синтетическая встреча','2026-09-23',ANALYSIS,SEGMENTS,NAMES,
                                include_transcript=False,source='synthetic_text'))
    assert 'Участники' in text
    assert 'Айжан Қали' in text
    assert 'Источник: синтетический демонстрационный текст.' in text
    assert 'Синтетическое предупреждение' in text
    assert 'Транскрипт' not in text


def test_docx_rederives_task_quote_from_segments():
    text = docx_text(build_docx('Синтетическая встреча','2026-09-23',ANALYSIS,SEGMENTS,NAMES,
                                include_transcript=False))
    assert '[S1] Синтетикалық тапсырма мәтіні.' in text
    assert 'Поддельная цитата' not in text


def test_docx_rejects_invalid_review_before_export():
    broken = {'summary':'Синтетическое саммари','tasks':[{'task':'Проверить','source_ids':['S404']}],
              'warnings':[]}
    with pytest.raises(ValueError, match='исходн'):
        build_docx('Синтетическая встреча','2026-09-23',broken,SEGMENTS,NAMES,include_transcript=False)
