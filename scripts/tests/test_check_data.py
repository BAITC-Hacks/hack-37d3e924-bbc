"""File checks use explicitly synthetic temporary audio and DOCX, never models."""

import json
import os
import subprocess
import sys
import wave
from pathlib import Path

import pytest
from docx import Document

from scripts import check_data


@pytest.fixture
def synthetic_data(tmp_path):
    source = tmp_path / "synthetic-data"
    source.mkdir()
    with wave.open(str(source / "synthetic.wav"), "wb") as audio:
        audio.setnchannels(2)
        audio.setsampwidth(2)
        audio.setframerate(8000)
        audio.writeframes(b"\0" * 8000 * 2 * 2)
    document = Document()
    document.add_paragraph("SYNTHETIC_PRIVATE_PARAGRAPH")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "SYNTHETIC_PRIVATE_CELL"
    document.save(source / "synthetic.docx")
    return source


def test_real_decode_docx_roundtrip_sources_and_report_privacy(
    synthetic_data, tmp_path
):
    before = {path.name: path.read_bytes() for path in synthetic_data.iterdir()}
    report, path = check_data.run_checks(
        synthetic_data, tmp_path / "reports", tmp_path / "no-models"
    )
    assert report["status"] == "PASS"
    audio = next(item for item in report["files"] if item["kind"] == "audio")
    document = next(item for item in report["files"] if item["kind"] == "docx")
    assert (audio["sample_rate"], audio["channels"], audio["duration_seconds"]) == (
        16000,
        1,
        1.0,
    )
    assert audio["diarization"] == {"status": "NOT_RUN"}
    assert document["checks"] == {"text_preserved": True, "has_content": True}
    assert (
        document["paragraph_count"],
        document["table_count"],
        document["cell_count"],
    ) == (1, 1, 1)
    assert all(item["source_unchanged"] for item in report["files"])
    assert before == {
        source.name: source.read_bytes() for source in synthetic_data.iterdir()
    }
    assert "SYNTHETIC_PRIVATE" not in path.read_text(encoding="utf-8")
    assert list(path.parent.iterdir()) == [path]
    _, second = check_data.run_checks(
        synthetic_data, path.parent.parent, tmp_path / "no-models"
    )
    assert second != path and path.is_file()


def test_requested_diarization_without_weights_is_partial(
    synthetic_data, tmp_path, monkeypatch
):
    monkeypatch.setattr(
        check_data, "diarize", lambda *_: pytest.fail("Missing models must not load")
    )
    report, _ = check_data.run_checks(
        synthetic_data, tmp_path / "reports", tmp_path / "missing", True
    )
    assert report["status"] == "PARTIAL"
    audio = next(item for item in report["files"] if item["kind"] == "audio")
    assert audio["diarization"]["status"] == "BLOCKED"


@pytest.mark.parametrize(
    "turns, expected",
    [
        ([{"start": 0.0, "end": 0.5, "speaker_id": "SPEAKER_00"}], "PASS"),
        ([{"start": 0.0, "end": 3.0, "speaker_id": "SPEAKER_00"}], "FAIL"),
        (
            [
                {"start": 0.5, "end": 0.8, "speaker_id": "SPEAKER_00"},
                {"start": 0.0, "end": 0.3, "speaker_id": "SPEAKER_01"},
            ],
            "FAIL",
        ),
        ([{"start": 0.0, "end": float("nan"), "speaker_id": "SPEAKER_00"}], "FAIL"),
        ([], "FAIL"),
    ],
)
def test_optional_diarization_validates_canonical_turns(
    synthetic_data, tmp_path, monkeypatch, turns, expected
):
    models = tmp_path / "synthetic-models"
    (models / "diarization").mkdir(parents=True)
    for filename in ("segmentation.onnx", "embedding.onnx"):
        (models / "diarization" / filename).write_bytes(
            b"synthetic placeholder; never loaded"
        )

    def fake_diarize(audio_path, settings):
        assert audio_path.is_file()
        assert settings.diarizer == "sherpa" and settings.device == "cpu"
        assert settings.diarization_path == str(models / "diarization")
        return {"turns": turns, "overlap": False}

    monkeypatch.setattr(check_data, "diarize", fake_diarize)
    report, path = check_data.run_checks(
        synthetic_data, tmp_path / "reports", models, True
    )
    assert report["status"] == expected
    assert "SPEAKER_00" not in path.read_text(encoding="utf-8")


def test_cli_sanitizes_invalid_inputs_and_no_content_escapes(tmp_path):
    source = tmp_path / "synthetic-invalid"
    source.mkdir()
    for name in ("invalid.wav", "invalid.docx"):
        (source / name).write_bytes(b"SYNTHETIC_PRIVATE_INVALID_CONTENT")
    result = subprocess.run(
        [
            sys.executable,
            str(Path(check_data.__file__)),
            "--data-dir",
            str(source),
            "--output-dir",
            str(tmp_path / "reports"),
        ],
        capture_output=True,
        text=True,
        timeout=30,
    )
    assert result.returncode == 1
    summary = json.loads(result.stdout)
    assert summary["status"] == "FAIL"
    report = Path(summary["report"]).read_text(encoding="utf-8")
    assert "SYNTHETIC_PRIVATE" not in result.stdout + result.stderr + report
    assert all(item["source_unchanged"] for item in json.loads(report)["files"])


def test_empty_directory_fails(tmp_path):
    source = tmp_path / "empty"
    source.mkdir()
    report, _ = check_data.run_checks(source, tmp_path / "reports", tmp_path / "models")
    assert report["status"] == "FAIL" and report["files"] == []


def test_network_guard_is_explicit_and_process_scoped():
    code = """import os, socket
from scripts.check_data import install_offline_guard
install_offline_guard()
assert os.environ['HF_HUB_OFFLINE'] == '1'
for action in (lambda: socket.getaddrinfo('example.com', 443),
               lambda: socket.socket().connect(('127.0.0.1', 9))):
    try:
        action()
    except PermissionError:
        pass
    else:
        raise AssertionError('Network guard did not block')
print('blocked')
"""
    result = subprocess.run(
        [sys.executable, "-c", code],
        cwd=check_data.ROOT,
        capture_output=True,
        text=True,
        timeout=15,
        env=dict(os.environ),
    )
    assert result.returncode == 0, result.stderr
    assert result.stdout.strip() == "blocked"
